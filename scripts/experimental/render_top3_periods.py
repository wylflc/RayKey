#!/usr/bin/env python3
"""Render the historical top-three CSV evidence as Chinese Excel and Markdown."""
import argparse
import csv
import json
import os
from collections import defaultdict
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
LABELS={
    'segment_id':'段号','segment_start':'本段开始','segment_end':'本段结束','trading_days':'持续交易日',
    'signal_date':'信号日','rank':'段首排名','security_code':'股票代码','security_name':'股票名称',
    'pv':'段首P/V','signal_close':'信号收盘价','intrinsic_value':'当时合理价V','ma20':'当日复权MA20',
    'ma60':'当日复权MA60','band_available_at':'估值输入可得日','eligible_count':'当日全部合格只数',
    'execution_date':'测量起点(T+1)','execution_close':'起点收盘价','price_history_end':'该股统计行情末日',
    'ma60_break_date':'MA60周期破位日','group':'样本口径','years':'自然年数','signals':'信号条数',
    'complete':'完整窗口数','missing':'缺失窗口数','signal_companies':'信号公司数',
    'complete_companies':'完整窗口公司数','median_return':'累计回报中位','mean_return':'累计回报均值',
    'positive_fraction':'正收益占比','max_company_fraction':'最大单公司样本占比','signal_year':'信号年份',
    'period':'年度/季度','first_quote':'首个报价日','last_quote':'源行情末日','cheap_states':'低PV状态条数',
    'top3_count':'前三名实际只数','activation_date':'首次激活日','anchor_date':'虚拟定锚日',
    'stop_anchor':'当日已调整止损锚','stop_line':'当日生效止损线','qualification':'资格来源',
    'activation_close':'激活日收盘','activation_ma20':'激活日MA20','activation_ma60':'激活日MA60',
    'original_anchor':'最初止损锚','last_adjusted_anchor':'末次已调整锚','reset_date':'止损复位日',
    'reset_close':'复位日收盘','reset_stop_line':'复位日止损线'}
for h in (1,3):
    LABELS.update({f'target_{h}y':f'{h}年自然周年',f'end_{h}y':f'{h}年测量终点',
                   f'status_{h}y':f'{h}年窗口状态',f'end_close_{h}y':f'{h}年终点收盘价',
                   f'return_{h}y':f'{h}年累计总回报'})
GROUPS={'daily_top3':'逐日前三','segment_starts':'名单分段段首','quarterly_top3':'季度首日',
        'annual_top3':'年度首日','ma60_episodes':'MA60周期去重','ma60_nonoverlap':'MA60周期+同股窗口不重叠',
        'activation_first_top3':'每个激活周期首次进入前三'}
STATUS={'complete':'完整','execution_after_cutoff':'起点超出行情末日',
        'execution_quote_missing':'T+1无报价','forward_incomplete':'尚未满期',
        'endpoint_after_stock_history':'个股行情提前终止','endpoint_quote_missing':'终点无报价'}


def read(path):
    with path.open(encoding='utf-8-sig',newline='') as f:
        return list(csv.DictReader(f))


def pct(value):
    return f'{float(value):+.2%}' if value not in ('',None) else '—'


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source',type=Path,required=True)
    ap.add_argument('--report',type=Path,required=True)
    args=ap.parse_args()
    source=args.source
    m=json.loads((source/'manifest.json').read_text())
    latched=m.get('eligibility')=='stop-latched'
    relative_source=Path(os.path.relpath(source.resolve(),args.report.parent.resolve())).as_posix()
    segments=read(source/'segments_top3.csv')
    summary=read(source/'summary.csv')
    grouped=defaultdict(list)
    for r in segments:grouped[int(r['segment_id'])].append(r)
    nonempty=sum(any(r['security_code'] for r in rs) for rs in grouped.values())
    notes=[
        f"范围：2009-11-02至{m['cutoff']}；{m['membership_segments']}个名单区段，其中{nonempty}段有合格股票。",
        '分段：前三名代码集合有人进入/退出即换段；仅内部换位不换段。每段P/V及排名取段首，结束日为事后识别；最后一段仅截至数据末日，尚不知真实结束日。',
        f"资格：当前BASE历史时点面板在册，0<P/V≤{m['buy_line']:.4f}，收盘>复权MA20>复权MA60。全体合格股票按P/V升序，同值按代码升序。",
        '收益：段首信号的下一市场交易日收盘买入，至其后1/3个自然周年遇休市顺延到首个市场交易日收盘；为累计总回报。',
        '总回报含现金分红再投、送转和配股认购成本，不含税费、融资或滑点。区段结束/跌破MA60不截断固定期收益；不是实际账户收益。',
        '缺失：T+1或周年终点停牌不推迟测量，尚未满期留空；最新日信号的T+1超出行情末端亦留空。Excel保留每条具体状态及起止日。',
        '行情：统一截至沪深300本地行情末日2026-08-07；部分个股源文件虽至08-28，不使用该段不齐全横截面。没有刷新行情。',
        '模型：用现行候选侧估值状态回看历史；不是逐年沿用当年旧版策略的操作记录。P/V来自源状态四位小数，V也有源舍入误差。',
        '披露时点沿用BASE：晚间报告按公告日戳前一个市场交易日生效；原可得日戳保留在明细中，本次未独立重建公告时刻。',
        '资格边界：回测没有人工冻结/review_pending和L3战术理由；未模拟账户资金、已有持仓、股债上限、单票上限、冷却及真实成交。',
        '同一股票可在多个区段重复出现；每日/段首信号行不是独立样本。MA60去重及同股窗口不重叠是附加描述，仍有跨股共同市场影响。',
        '历史面板存在回溯判断、财报修订及终点删失限制；收益统计只使用完整窗口，缺窗分列。名称用于辨认公司，代码为主键。',
        '核验：见 verification.json；完整复现配置、输入/输出哈希及作业编号见 manifest.json。']
    if latched:
        notes[2]=f"资格：历史面板在册且0<P/V≤{m['buy_line']:.4f}；首次收盘>MA20>MA60激活，此后只需MA20>MA60。按当日P/V升序，同值按代码升序。"
        notes[9]='激活不要求真实成交或已进入前三。MA20≤MA60、P/V超线或暂离面板只暂停当天排名，不清除记忆；首次激活从研究起点开始。人工冻结/L3战术、资金仓位等账户闸门未模拟。'
        notes.insert(3,'唯一复位：收盘跌破min(虚拟建仓MA60锚,当日MA60)。锚取首次信号后下一报价日MA60，随后按分红送配调整；T+1停牌时仅虚拟锚顺延，收益起点仍按固定市场T+1、缺价留空。信号当日尚未定锚，不提前使用未来均线。')
        notes.insert(4,'资格周期与MA60统计去重分开：跌破当日MA60但未跌破生效止损线不清除激活。Excel另列全部激活/复位周期，以及每周期首次进入前三的收益。')
        notes.insert(5,'MA60辅助计数保留已经低于MA60的合格信号，记作当日即结束的零时长观测，可能连续出现；不能解释为新的激活或建仓。判断资格持续性应看激活/止损周期，固定期收益仍完整保留。')
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.formatting.rule import CellIsRule
    wb=Workbook()
    ws=wb.active;ws.title='口径说明';ws.append(['项目','说明'])
    for i,note in enumerate(notes,1):ws.append([i,note])
    ws.column_dimensions['A'].width=12;ws.column_dimensions['B'].width=130
    for row in ws.iter_rows(min_row=2):
        row[1].alignment=Alignment(wrap_text=True,vertical='top');ws.row_dimensions[row[0].row].height=42
    sheets=[('名单变化分段','segments_top3.csv'),('逐日完整明细','daily_top3.csv'),
            ('汇总及去重','summary.csv'),('MA60周期','ma60_episodes.csv'),
            ('同股1年不重叠','ma60_nonoverlap_1y.csv'),('同股3年不重叠','ma60_nonoverlap_3y.csv'),
            ('逐年分布','by_year.csv'),('逐公司分布','by_company.csv'),
            ('年度首日辅助','annual_top3.csv'),('季度首日辅助','quarterly_top3.csv'),
            ('逐日资格覆盖','daily_coverage.csv'),('源行情覆盖','coverage.csv')]
    if latched:
        sheets[3:3]=[('激活及止损复位','activation_cycles.csv'),('激活周期首次前三','activation_first_top3.csv')]
    integer={'segment_id','rank','trading_days','eligible_count','years','signals','complete','missing',
             'signal_companies','complete_companies','cheap_states','top3_count'}
    numeric={'pv','signal_close','intrinsic_value','ma20','ma60','execution_close',
             'end_close_1y','end_close_3y','stop_anchor','stop_line','activation_close','activation_ma20',
             'activation_ma60','original_anchor','last_adjusted_anchor','reset_close','reset_stop_line'}
    percent={'return_1y','return_3y','median_return','mean_return','positive_fraction','max_company_fraction'}
    from openpyxl.utils import get_column_letter
    for title,filename in sheets:
        rows=read(source/filename);fields=list(rows[0]);ws=wb.create_sheet(title)
        # Put the requested figures together; detailed audit dates remain to the right.
        leading=[k for k in ('segment_id','segment_start','segment_end','trading_days','signal_date',
                            'rank','security_code','security_name','pv','return_1y','return_3y') if k in fields]
        fields=leading+[k for k in fields if k not in leading]
        ws.append([LABELS.get(k,k) if filename=='segments_top3.csv' else
                   {'rank':'当时排名','pv':'当时P/V'}.get(k,LABELS.get(k,k)) for k in fields])
        for row in rows:
            values=[]
            for k in fields:
                v=row[k]
                if not v:v=None
                elif k in integer:v=int(v)
                elif k in numeric|percent:v=float(v)
                elif k=='group':v=GROUPS.get(v,v)
                elif k=='qualification':v={'initial':'首次触发','continued':'激活后延续','continued_below_ma20':'回踩MA20仍合格'}.get(v,v)
                elif k.startswith('status_'):v=STATUS.get(v,v)
                values.append(v)
            ws.append(values)
        ws.auto_filter.ref=ws.dimensions;ws.freeze_panes='I2' if filename=='segments_top3.csv' else 'A2'
        for col,k in enumerate(fields,1):
            letter=get_column_letter(col)
            ws.column_dimensions[letter].width=18 if k!='group' else 32
            if k in percent:
                for cells in ws.iter_rows(min_row=2,min_col=col,max_col=col):cells[0].number_format='+0.00%;-0.00%;0.00%'
                region=f'{letter}2:{letter}{ws.max_row}'
                ws.conditional_formatting.add(region,CellIsRule(operator='lessThan',formula=['0'],font=Font(color='008044')))
                ws.conditional_formatting.add(region,CellIsRule(operator='greaterThan',formula=['0'],font=Font(color='C32626')))
            elif k in numeric:
                for cells in ws.iter_rows(min_row=2,min_col=col,max_col=col):cells[0].number_format='0.0000'
    for ws in wb:
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(color='FFFFFF',bold=True)
        ws.row_dimensions[1].height=30
    wb.save(source/'历史前三名分段及后续收益.xlsx')
    report=['# 历史买入合格集：P/V前三名名单分段与后续收益','',
            f"**{m['membership_segments']}个名单区段，{m['market_days']}个市场交易日，{m['results'][0]['signals']}条逐日前三信号。**",'',
            '## 统计口径','']+['- '+n for n in notes]+['','## 汇总核对','',
            '以下为单只股票累计回报的中位数，不能当作组合年化收益。各期内相同股票会重复计数。','',
            '| 样本 | 信号数 | 1年完整/公司 | 1年中位 | 3年完整/公司 | 3年中位 |',
            '| --- | ---: | ---: | ---: | ---: | ---: |']
    if latched:report[0]='# 历史P/V前三名：首次信号激活、止损后复位'
    for g in ('segment_starts','daily_top3','ma60_episodes','ma60_nonoverlap') + (('activation_first_top3',) if latched else ()):
        a=next(r for r in summary if r['group']==g and r['years']=='1')
        b=next(r for r in summary if r['group']==g and r['years']=='3')
        n=a['signals'] if a['signals']==b['signals'] else f"{a['signals']}/{b['signals']}"
        report.append(f"| {GROUPS[g]} | {n} | {a['complete']}/{a['complete_companies']} | {pct(a['median_return'])} | {b['complete']}/{b['complete_companies']} | {pct(b['median_return'])} |")
    if latched:
        c=m['eligibility_comparison']
        report+=['','## 对上一版的更正','',
                 '上一版每天重新要求收盘高于MA20，未记录首次触发后的资格保持。本版采用用户确认的信号激活、止损复位规则。','',
                 '| 指标 | 旧版每天严格趋势 | 本版激活后延续 |','| --- | ---: | ---: |',
                 f"| 名单区段数 | {c['strict_segments']} | {c['latched_segments']} |",
                 f"| 每段平均交易日 | {c['strict_average_days']:.2f} | {c['latched_average_days']:.2f} |",
                 f"| 每段交易日中位 | {c['strict_median_days']:.1f} | {c['latched_median_days']:.1f} |",
                 f"| 仅持续1日的段数 | {c['strict_one_day_segments']} | {c['latched_one_day_segments']} |",'',
                 f"{c['changed_membership_days']}个交易日的前三成员不同；本版有{c['below_ma20_top3_rows']}条前三信号来自已激活股票回踩MA20。激活共{c['activations']}次，止损复位{c['stop_resets']}次。名单仍会因相对P/V排序、估值更新、MA20/MA60排列及止损变化；未额外人为合并这些变化。"]
    report+=['','## 全部分段明细','',
             '每格顺序为 **股票（代码） · P/V · 后1年 · 后3年**。排名固定为段首排名。`未满`为行情尚未覆盖完整周年，其余缺失明确列因。','',
             f'完整日期、价格、缺失原因与逐日变化见[Excel]({relative_source}/历史前三名分段及后续收益.xlsx)及[分段CSV]({relative_source}/segments_top3.csv)。']
    previous_year=''
    def cell(r):
        if not r['security_code']:return '无合格标的'
        def ret(h):
            return pct(r[f'return_{h}y']) if r[f'return_{h}y'] else ('未满' if r[f'status_{h}y']=='forward_incomplete' else STATUS.get(r[f'status_{h}y'],'—'))
        return f"{r['security_name']}（{r['security_code']}）<br>{float(r['pv']):.4f} · {ret(1)} · {ret(3)}"
    for ident,rows in grouped.items():
        r=rows[0];year=r['segment_start'][:4]
        if year!=previous_year:
            report+=['',f'### {year}年开始的区段','',
                     '| 段号 | 起止日期（交易日数） | 第1名：P/V · 1年 · 3年 | 第2名：P/V · 1年 · 3年 | 第3名：P/V · 1年 · 3年 |',
                     '| ---: | --- | --- | --- | --- |']
            previous_year=year
        report.append(f"| {ident} | {r['segment_start']}～{r['segment_end']}（{r['trading_days']}） | "+' | '.join(cell(x) for x in rows)+' |')
    verification=json.loads((source/'verification.json').read_text())
    report+=['','## 复现与验证','',
             f"统计作业 {m['job_id']}；输入流式读取{m['states_rows_read']:,}行。已核对全部{verification['daily_rows_checked']:,}条日信号和{verification['segments_checked']:,}个区段，独立重算{verification['full_universe_rank_dates']}个日期的全池排名、{verification['direct_ma_windows']:,}个均线窗口及{verification['independent_cash_ledgers']}条现金台账；收益最大误差{verification['max_return_error']:.2g}。边界单元检查及既有周期/除权检查合计{23 if latched else 14}项通过。",'',
             ('复现入口：`scripts/slurm/top3_latched_20260915.sbatch`；完成后运行 `scripts/slurm/top3_latched_verify_20260915.sbatch`。' if latched else
              '复现入口：`scripts/slurm/top3_periods_20260915.sbatch`；完成后运行 `scripts/slurm/top3_verify_20260915.sbatch`。')+
             f'后者先独立核验，再生成本报告与Excel。输入方案见[预登记]({relative_source}/preregister.md)，哈希见同目录 `manifest.json`。']
    args.report.write_text('\n'.join(report)+'\n')
    print(json.dumps({'segments':len(grouped),'nonempty_segments':nonempty,'excel':str(source/'历史前三名分段及后续收益.xlsx'),'report':str(args.report)},ensure_ascii=False))


if __name__=='__main__':main()
